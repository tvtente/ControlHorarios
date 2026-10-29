"""Completa los descansos de prueba y ordena cronológicamente los fichajes existentes."""

from collections import defaultdict
from copy import copy
from datetime import date, timedelta
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from excel_manager import ExcelManager


ACCIONES = ("Entrada a Turno", "Inicio de Descanso", "Fin de Descanso", "Salida de Turno")
DATOS_PRUEBA = "Prueba: jornada histórica"


def ajustar(archivo):
    archivo = Path(archivo)
    libro = load_workbook(archivo)
    try:
        hoja = libro["Registros"]
        encabezados = [celda.value for celda in hoja[1]]
        registros = [dict(zip(encabezados, fila)) for fila in hoja.iter_rows(min_row=2, values_only=True)
                     if fila[0] is not None]
        grupos = defaultdict(list)
        for registro in registros:
            grupos[(registro["id_empleado"], str(registro["fecha"]))].append(registro)
        fechas_ocupadas = {clave for clave in grupos}
        siguiente_id = max(r["id_registro"] for r in registros) + 1
        descansos_nuevos = 0
        jornadas_movidas = 0

        for (empleado, fecha), grupo in list(grupos.items()):
            prueba = [r for r in grupo if r["observaciones"] == DATOS_PRUEBA]
            if [r["accion"] for r in prueba] != ["Entrada a Turno", "Salida de Turno"]:
                continue
            entrada, salida = prueba
            if len(grupo) > len(prueba):
                dia = date.fromisoformat(fecha) - timedelta(days=1)
                while (empleado, dia.isoformat()) in fechas_ocupadas:
                    dia -= timedelta(days=1)
                fechas_ocupadas.add((empleado, dia.isoformat()))
                for registro in prueba:
                    registro["fecha"] = dia.isoformat()
                entrada["hora"] = entrada["hora_inicio_prevista"] + ":00"
                salida["hora"] = salida["hora_fin_prevista"] + ":00"
                jornadas_movidas += 1

            inicio = entrada["hora_inicio_descanso_prevista"] + ":00"
            fin = entrada["hora_fin_descanso_prevista"] + ":00"
            if not entrada["hora"] < inicio < fin < salida["hora"]:
                raise ValueError(f"Horario incoherente en {empleado}, {entrada['fecha']}.")
            for accion, hora in ((ACCIONES[1], inicio), (ACCIONES[2], fin)):
                nuevo = {nombre: entrada.get(nombre) for nombre in encabezados[:12]}
                nuevo.update(id_registro=siguiente_id, accion=accion, hora=hora)
                registros.append(nuevo)
                siguiente_id += 1
                descansos_nuevos += 1

        registros.sort(key=lambda r: (str(r["fecha"]), str(r["hora"]),
                                     str(r["id_empleado"]), ACCIONES.index(r["accion"])))
        verificar_orden(registros)

        formulas = {
            nombre: (hoja.cell(2, col).value, hoja.cell(2, col).coordinate)
            for col, nombre in enumerate(encabezados, 1)
            if nombre in ("diferencia_min", "cumplimiento")
        }
        estilos = [copy(celda._style) for celda in hoja[2]]
        hoja.delete_rows(2, hoja.max_row - 1)
        for fila, registro in enumerate(registros, 2):
            for columna, nombre in enumerate(encabezados, 1):
                celda = hoja.cell(fila, columna)
                celda._style = copy(estilos[columna - 1])
                if nombre in encabezados[:12]:
                    celda.value = registro.get(nombre)
                elif nombre in formulas and formulas[nombre][0]:
                    formula, origen = formulas[nombre]
                    celda.value = Translator(formula, origin=origen).translate_formula(celda.coordinate)
                else:
                    celda.value = None
        referencia = f"A1:{get_column_letter(len(encabezados))}{len(registros) + 1}"
        for tabla in hoja.tables.values():
            tabla.ref = referencia
            if tabla.autoFilter is not None:
                tabla.autoFilter.ref = referencia
        libro.save(archivo)
    finally:
        libro.close()

    ExcelManager(archivo)
    return len(registros), descansos_nuevos, jornadas_movidas


def verificar_orden(registros):
    """Comprueba que cada jornada terminada tenga cuatro horas crecientes."""
    grupos = defaultdict(list)
    for registro in registros:
        grupos[(registro["id_empleado"], str(registro["fecha"]))].append(registro)
    for clave, grupo in grupos.items():
        acciones = [r["accion"] for r in grupo]
        if acciones == ["Entrada a Turno"]:
            continue
        if acciones != list(ACCIONES):
            raise ValueError(f"Secuencia incompleta o duplicada en {clave}: {acciones}")
        horas = [str(r["hora"]) for r in grupo]
        if horas != sorted(horas) or len(set(horas)) != 4:
            raise ValueError(f"Las horas no son crecientes en {clave}: {horas}")


if __name__ == "__main__":
    ruta = Path(__file__).resolve().parent / "data" / "control_horarios.xlsx"
    total, descansos, movidas = ajustar(ruta)
    print(f"Fichajes: {total}; descansos añadidos: {descansos}; jornadas movidas: {movidas}")
