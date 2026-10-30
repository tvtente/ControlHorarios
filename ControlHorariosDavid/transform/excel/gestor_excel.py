"""Lectura y escritura del libro de fichajes, sin dependencias de la interfaz."""

from pathlib import Path

import openpyxl as pxl

from config import NOMBRE_EXCEL, COLUMNAS_EXCEL

RUTA = Path(__file__).resolve().parent.parent


class GestorExcel:
    def __init__(self, ruta=RUTA, nombre_excel=NOMBRE_EXCEL):
        self.ruta = Path(ruta)
        self.nombre_excel = nombre_excel
        self.ruta_excel = self.ruta / self.nombre_excel

    def inicializar_excel(self):
        """Abre el libro existente o crea uno con sus columnas originales."""
        if self.ruta_excel.exists():
            libro = pxl.load_workbook(self.ruta_excel)
        else:
            self.ruta.mkdir(parents=True, exist_ok=True)
            libro = pxl.Workbook()
            libro.active.title = "Registros"
            libro.active.append(COLUMNAS_EXCEL)
            try:
                libro.save(self.ruta_excel)
            except Exception:
                libro.close()
                raise
        return libro, libro.active

    def obtener_registros_excel(self):
        libro, hoja = self.inicializar_excel()
        try:
            return list(hoja.iter_rows(min_row=2, values_only=True))
        finally:
            libro.close()

    def guardar_excel(self, registro):
        libro, hoja = self.inicializar_excel()
        try:
            hoja.append(registro)
            libro.save(self.ruta_excel)
        finally:
            libro.close()
