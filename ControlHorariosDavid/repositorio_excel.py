"""Persistencia de fichajes en el libro Excel del proyecto de David."""

from pathlib import Path

from openpyxl import Workbook, load_workbook


class RepositorioExcel:
    COLUMNAS = ("ID", "Nombre", "Departamento", "Acción", "Fecha", "Hora", "Observaciones")

    def __init__(self, ruta):
        self.ruta = Path(ruta)

    def inicializar_excel(self):
        if self.ruta.exists():
            libro = load_workbook(self.ruta)
        else:
            libro = Workbook()
            libro.active.append(self.COLUMNAS)
            libro.save(self.ruta)
        return libro, libro.active

    def obtener_registros_excel(self):
        libro, hoja = self.inicializar_excel()
        try:
            return list(hoja.iter_rows(min_row=2, values_only=True))
        finally:
            libro.close()

    def guardar_excel(self, registro):
        libro, hoja = self.inicializar_excel()
        try:
            hoja.append(registro)
            libro.save(self.ruta)
        finally:
            libro.close()
