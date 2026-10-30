import openpyxl as pxl
from pathlib import Path
from config import  NOMBRE_EXCEL, COLUMNAS_EXCEL

RUTA = Path(__file__).resolve().parent.parent

class GestorExcel:
    def __init__(self, ruta=RUTA, nombre_excel=NOMBRE_EXCEL):
        self.ruta = Path(ruta)
        self.nombre_excel = nombre_excel
        self.archivo_excel = self.ruta / self.nombre_excel
        self.ruta_excel = Path(self.archivo_excel)

    def inicializar_excel(self):
        """Inicializa el archivo de Excel y devuelve el libro y la hoja activa."""
        if not self.ruta_excel.exists():
            libro = pxl.Workbook()
            hoja = libro.active
            hoja.title = "Registros"
            hoja.append(COLUMNAS_EXCEL)
            try:
                libro.save(self.ruta_excel)
            except Exception:
                libro.close()
                raise 
            return libro, hoja
        
        libro = pxl.load_workbook(self.ruta_excel)
        hoja = libro.active
        return libro, hoja
        
    def obtener_registros_excel(self):
        """Obtenemos todos los registros del Excel."""
        libro = None
        try:
            libro, hoja = self.inicializar_excel()
            registros = list(hoja.iter_rows(min_row=2, values_only=True))
            return registros
        finally:
            if libro is not None:
                libro.close()
                
    def guardar_excel(self, registro):
        libro = None
        try:
            libro, hoja = self.inicializar_excel()
            hoja.append(registro)
            libro.save(self.ruta_excel)
        finally:
            if libro is not None:
                libro.close()
                
