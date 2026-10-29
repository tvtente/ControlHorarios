import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_manager import ExcelManager


class HorasTrabajadasTest(unittest.TestCase):
    def test_migracion_conserva_fichajes_y_descuenta_descanso(self):
        with tempfile.TemporaryDirectory() as carpeta:
            archivo = Path(carpeta) / "control.xlsx"
            libro = Workbook()
            hoja = libro.active
            hoja.title = "Registros"
            hoja.append([
                "id_registro", "id_empleado", "nombre_completo", "departamento",
                "accion", "fecha", "hora", "observaciones", "hora_inicio_prevista",
                "hora_fin_prevista", "diferencia_min", "cumplimiento",
                "hora_inicio_descanso_prevista", "hora_fin_descanso_prevista",
            ])
            for numero, (accion, hora) in enumerate((
                ("Entrada a Turno", "08:00:00"),
                ("Inicio de Descanso", "12:00:00"),
                ("Fin de Descanso", "12:30:00"),
                ("Salida de Turno", "17:00:00"),
            ), 1):
                hoja.append([numero, "EMP-001", "Ana", "Ventas", accion,
                             "2026-10-28", hora, "", "08:00", "17:00", None,
                             None, "12:00", "12:30"])
            libro.save(archivo)
            libro.close()

            gestor = ExcelManager(archivo)
            registros = gestor.consultar_registros()
            self.assertEqual(len(registros), 4)
            self.assertEqual(registros[-1]["horas_trabajadas"], "8:30")
            self.assertEqual(registros[-1]["cumplimiento"], "Cumple horario")
            libro = load_workbook(archivo)
            hoja = libro["Registros"]
            self.assertEqual(
                [hoja.cell(1, columna).value for columna in range(9, 13)],
                ["hora_inicio_prevista", "hora_inicio_descanso_prevista",
                 "hora_fin_descanso_prevista", "hora_fin_prevista"],
            )
            self.assertEqual(hoja["E5"].value, "Salida de Turno")
            self.assertEqual(hoja["P5"].value.strftime("%H:%M"), "12:00")
            self.assertEqual(hoja["Q5"].value.strftime("%H:%M"), "12:30")
            self.assertEqual(
                hoja["S5"].value,
                '=IF(OR(O5="",R5=""),"",MOD(R5-O5,1)-IF(OR(P5="",Q5=""),0,MOD(Q5-P5,1)))',
            )
            self.assertIsNone(hoja["S4"].value)
            libro.close()
            calculado = load_workbook(archivo, data_only=True)
            self.assertEqual(calculado["Registros"]["S5"].value.total_seconds(), 8.5 * 3600)
            calculado.close()

            # El histórico toma el valor guardado por Excel, sin recalcular la duración.
            gestor._guardar_cache_horas({"S5": 6 / 24})
            self.assertEqual(gestor.consultar_registros()[-1]["horas_trabajadas"], "6:00")


if __name__ == "__main__":
    unittest.main()
