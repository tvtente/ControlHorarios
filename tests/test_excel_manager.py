import tempfile
import unittest
from datetime import time
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_manager import ExcelManager


class HorariosTests(unittest.TestCase):
    def test_resuelve_horarios_sin_modificar_formulas(self):
        with tempfile.TemporaryDirectory() as carpeta:
            archivo = Path(carpeta) / 'horarios.xlsx'
            libro = Workbook()
            hoja = libro.active
            hoja.title = 'Registros'
            hoja.append(ExcelManager.HEADERS)
            formula = '=IFERROR(VLOOKUP(D2,Departamentos!$A$2:$D$100,2,FALSE),"")'
            fin = '=IFERROR(VLOOKUP(D2,Departamentos!$A$2:$D$100,3,FALSE),"")'
            hoja.append([1, 'A', 'Ana', 'Reservas', 'Entrada', '2026-10-28',
                         '08:30', '', formula, fin])
            hoja.append([2, 'B', 'Bea', 'Desconocido', 'Entrada', '2026-10-28',
                         '09:00', '', formula, fin])
            hoja.append([3, 'C', 'Cris', 'Reservas', 'Entrada', '2026-10-28',
                         '09:00', '', '10:00', '18:00'])
            # Las filas con fórmulas pero sin fichaje no son registros.
            hoja.append([None] * 8 + [formula, fin])
            departamentos = libro.create_sheet('Departamentos')
            departamentos.append(['Departamento', 'Hora entrada', 'Hora salida'])
            departamentos.append(['Reservas', time(8, 30), 17 / 24])
            libro.save(archivo)
            libro.close()
            registros = ExcelManager(archivo).consultar_registros()
            self.assertEqual(len(registros), 3)
            self.assertEqual(registros[0]['hora_inicio_prevista'], '08:30')
            self.assertEqual(registros[0]['hora_fin_prevista'], '17:00')
            self.assertIsNone(registros[1]['hora_inicio_prevista'])
            self.assertEqual(registros[2]['hora_inicio_prevista'], '10:00')
            libro = load_workbook(archivo)
            self.assertEqual(libro['Registros']['I2'].value, formula)
            libro.close()

    def test_horario_se_consulta_de_nuevo_y_admite_columnas_reordenadas(self):
        with tempfile.TemporaryDirectory() as carpeta:
            archivo = Path(carpeta) / 'horarios.xlsx'
            manager = ExcelManager(archivo)
            libro = load_workbook(archivo)
            hoja = libro.create_sheet('Departamentos')
            hoja.append(['Hora salida', 'Departamento', 'Hora entrada'])
            hoja.append([17 / 24, 'Reservas', time(8, 30)])
            hoja.append(['16:30', 'Operaciones', '8:00'])
            libro.save(archivo)
            libro.close()
            self.assertEqual(manager.consultar_horario_departamento('Reservas'), {
                'hora_inicio_prevista': '08:30', 'hora_fin_prevista': '17:00'})
            self.assertEqual(manager.consultar_horario_departamento('Operaciones'), {
                'hora_inicio_prevista': '08:00', 'hora_fin_prevista': '16:30'})
            self.assertEqual(manager.consultar_horario_departamento('Desconocido'), {
                'hora_inicio_prevista': '', 'hora_fin_prevista': ''})
            libro = load_workbook(archivo)
            libro['Departamentos']['C2'] = '09:15'
            libro.save(archivo)
            libro.close()
            self.assertEqual(manager.consultar_horario_departamento('Reservas')
                             ['hora_inicio_prevista'], '09:15')
