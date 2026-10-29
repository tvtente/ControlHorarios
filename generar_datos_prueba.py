"""Sustituye los fichajes por ejemplos de turnos y descansos."""
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path
import shutil

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from excel_manager import ExcelManager

ARCHIVO = Path(__file__).resolve().parent / 'data' / 'control_horarios.xlsx'
EMPLEADOS = (
    ('Ana García', 'Reservas y Emisión', 4),
    ('Carlos López', 'Operaciones', 4),
    ('Lucía Martín', 'Recursos Humanos', 4),
    ('Pablo Torres', 'Administración', 4),
    ('Elena Ruiz', 'Atención al Cliente', 4),
    ('Diego Sánchez', 'Reservas y Emisión', 4),
    ('Sofía Romero', 'Operaciones', 4),
    ('Javier Díaz', 'Recursos Humanos', 4),
    ('Carmen Vega', 'Administración', 3),
    ('Andrés Pérez', 'Atención al Cliente', 2),
    ('Nuria Castro', 'Reservas y Emisión', 2),
    ('Marta Molina', 'Operaciones', 1),
)
ACCIONES = ('Entrada a Turno', 'Inicio de Descanso', 'Fin de Descanso', 'Salida de Turno')


def generar(archivo=ARCHIVO):
    archivo = Path(archivo)
    manager = ExcelManager.__new__(ExcelManager)
    manager.filename = archivo
    horarios = {d: manager.consultar_horario_departamento(d) for _, d, _ in EMPLEADOS}
    ahora = datetime.now().replace(microsecond=0)
    hoy = ahora.date()
    registros = []
    # 25 jornadas completas distribuidas entre hoy y hace 29 días.
    for i in range(25):
        dia = hoy - timedelta(days=29 - round(i * 29 / 24))
        nombre, departamento, _ = EMPLEADOS[i % len(EMPLEADOS)]
        horario = horarios[departamento]
        entrada = datetime.combine(dia, datetime.strptime(horario['hora_inicio_prevista'], '%H:%M').time())
        salida = datetime.combine(dia, datetime.strptime(horario['hora_fin_prevista'], '%H:%M').time())
        entrada += timedelta(minutes=(0, 5, -3, 12, 0)[i % 5])
        salida += timedelta(minutes=(0, 10, -5)[i % 3])
        if dia == hoy:
            salida = min(salida, ahora - timedelta(minutes=5))
            entrada = min(entrada, salida - timedelta(hours=1))
        movimientos = [('Entrada a Turno', entrada), ('Salida de Turno', salida)]
        if dia < hoy and i % 3 == 0:
            descanso_inicio = datetime.combine(
                dia, datetime.strptime(horario['hora_inicio_descanso_prevista'], '%H:%M').time()
            )
            descanso_fin = datetime.combine(
                dia, datetime.strptime(horario['hora_fin_descanso_prevista'], '%H:%M').time()
            )
            movimientos.extend([
                ('Inicio de Descanso', descanso_inicio),
                ('Fin de Descanso', descanso_fin),
            ])
        for accion, momento in movimientos:
            registros.append(dict(
                id_empleado=f'EMP-{i % len(EMPLEADOS) + 1:03d}', nombre_completo=nombre,
                departamento=departamento, accion=accion,
                fecha=momento.strftime('%Y-%m-%d'), hora=momento.strftime('%H:%M:%S'),
                observaciones='Prueba: jornada histórica', **horario,
            ))
    # Los diez últimos son entradas iniciales de mañana para la presentación.
    horario_idiomas = manager.consultar_horario_departamento('Idiomas')
    if not all(horario_idiomas.values()):
        raise ValueError('Falta configurar el departamento Idiomas en Excel.')
    nombres = ('Laura Fernández', 'Miguel Santos', 'Paula Navarro', 'Daniel Ortega',
               'Irene Vidal', 'Álvaro Gil', 'Clara León', 'Sergio Ramos', 'Eva Molina', 'Hugo Serrano')
    for i, nombre in enumerate(nombres, 13):
        registros.append(dict(
            id_empleado=f'EMP-{i:03d}', nombre_completo=nombre, departamento='Idiomas',
            accion='Entrada a Turno', fecha=(hoy + timedelta(days=1)).isoformat(),
            hora='08:00:00', observaciones='Prueba inicial: presentación de mañana',
            **horario_idiomas,
        ))
    registros.sort(key=lambda r: (r['fecha'], r['hora'], r['id_empleado']))
    assert len(registros) == 76
    backup = archivo.parent / 'backups' / f'{archivo.stem}_antes_pruebas60_{ahora:%Y%m%d_%H%M%S}.xlsx'
    backup.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(archivo, backup)
    libro = load_workbook(archivo)
    try:
        hoja = libro['Registros']
        headers = [c.value for c in hoja[1]]
        plantillas = [(c.value, c.coordinate, copy(c._style)) for c in hoja[2]]
        if hoja.max_row > 1:
            hoja.delete_rows(2, hoja.max_row - 1)
        for numero, registro in enumerate(registros, 1):
            registro['id_registro'] = numero
            for columna, header in enumerate(headers, 1):
                celda = hoja.cell(numero + 1, columna)
                valor, origen, estilo = plantillas[columna - 1]
                celda._style = copy(estilo)
                if header in registro:
                    celda.value = registro[header]
                elif isinstance(valor, str) and valor.startswith('='):
                    celda.value = Translator(valor, origin=origen).translate_formula(celda.coordinate)
        ref = f'A1:{get_column_letter(len(headers))}{len(registros) + 1}'
        for tabla in hoja.tables.values():
            tabla.ref = ref
            if tabla.autoFilter is not None:
                tabla.autoFilter.ref = ref
        libro.save(archivo)
    finally:
        libro.close()
    return len(registros)


if __name__ == '__main__':
    print(f'Fichajes de prueba preparados: {generar()}')
