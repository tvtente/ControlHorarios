from employee_status import validar_fichaje
from pathlib import Path
from datetime import datetime, time
import re
from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn
from copy import copy


class ExcelManager:

    HEADERS = [
        "id_registro",
        "id_empleado",
        "nombre_completo",
        "departamento",
        "accion",
        "fecha",
        "hora",
        "observaciones",
        "hora_inicio_prevista",
        "hora_inicio_descanso_prevista",
        "hora_fin_descanso_prevista",
        "hora_fin_prevista",
        "diferencia_min",
        "cumplimiento",
        "entrada_real",
        "inicio_descanso_real",
        "fin_descanso_real",
        "salida_real",
        "horas_trabajadas",
    ]

    ACCIONES_JORNADA = {
        "Entrada a Turno": "entrada_real",
        "Inicio de Descanso": "inicio_descanso_real",
        "Fin de Descanso": "fin_descanso_real",
        "Salida de Turno": "salida_real",
    }

    def __init__(self, filename="control_horarios.xlsx"):
        self.filename = Path(filename)
        self.inicializar_excel()

    # Crea el libro si no existe y conserva los registros anteriores.
    def inicializar_excel(self):
        if not self.filename.exists():
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Registros"
            sheet.append(self.HEADERS)
            workbook.save(self.filename)
            workbook.close()
            return

        workbook = load_workbook(self.filename)

        try:
            if "Registros" not in workbook.sheetnames:
                raise ValueError("El archivo de Excel no contiene la hoja «Registros».")

            sheet = workbook["Registros"]

            existing_headers = [
                sheet.cell(row=1, column=col).value
                for col in range(1, sheet.max_column + 1)
            ]

            # Reúne las horas previstas antes de las columnas calculadas.
            if not existing_headers or all(h is None for h in existing_headers):
                for indice, header in enumerate(self.HEADERS, 1):
                    sheet.cell(1, indice).value = header
            else:
                anteriores = {nombre: indice + 1 for indice, nombre in enumerate(existing_headers)}
                orden = self.HEADERS + [h for h in existing_headers if h not in self.HEADERS]
                if existing_headers != orden:
                    filas = [
                        [(sheet.cell(fila, col).value, copy(sheet.cell(fila, col)._style))
                         for col in range(1, sheet.max_column + 1)]
                        for fila in range(2, sheet.max_row + 1)
                    ]
                    for col, header in enumerate(orden, 1):
                        sheet.cell(1, col).value = header
                    for fila, datos in enumerate(filas, 2):
                        for col, header in enumerate(orden, 1):
                            celda = sheet.cell(fila, col)
                            anterior = anteriores.get(header)
                            celda.value = datos[anterior - 1][0] if anterior else None
                            if anterior:
                                celda._style = copy(datos[anterior - 1][1])
                            if celda.data_type == "f":
                                if header == "diferencia_min":
                                    celda.value = re.sub(rf"\bJ{fila}\b", f"L{fila}", celda.value)
                                elif header == "cumplimiento":
                                    celda.value = re.sub(rf"\bK{fila}\b", f"M{fila}", celda.value)
                    existing_headers = orden

            self._actualizar_jornadas(sheet)

            for tabla in sheet.tables.values():
                inicio_col, inicio_fila, fin_col, fin_fila = range_boundaries(tabla.ref)
                if (inicio_col, inicio_fila) != (1, 1):
                    continue
                nuevo_fin = len(existing_headers)
                tabla.tableColumns = [
                    TableColumn(id=columna, name=str(existing_headers[columna - 1]))
                    for columna in range(1, nuevo_fin + 1)
                ]
                tabla.ref = f"A1:{get_column_letter(nuevo_fin)}{fin_fila}"
                if tabla.autoFilter is not None:
                    tabla.autoFilter.ref = tabla.ref

            workbook.save(self.filename)
        finally:
            workbook.close()

    @classmethod
    def _actualizar_jornadas(cls, sheet):
        """Deja las horas reales y el tiempo neto en la fila de salida."""
        columnas = {sheet.cell(1, col).value: col for col in range(1, sheet.max_column + 1)}
        jornadas = {}
        for fila in range(2, sheet.max_row + 1):
            empleado = sheet.cell(fila, columnas["id_empleado"]).value
            fecha = sheet.cell(fila, columnas["fecha"]).value
            accion = sheet.cell(fila, columnas["accion"]).value
            hora = sheet.cell(fila, columnas["hora"]).value
            if not all((empleado, fecha, accion, hora)):
                continue
            clave = empleado
            if accion == "Entrada a Turno":
                jornadas[clave] = {}
            jornada = jornadas.setdefault(clave, {})
            campo = cls.ACCIONES_JORNADA.get(accion)
            if campo:
                jornada[campo] = hora
            if accion != "Salida de Turno":
                for nombre in (*cls.ACCIONES_JORNADA.values(), "horas_trabajadas"):
                    sheet.cell(fila, columnas[nombre]).value = None
                continue
            for nombre in cls.ACCIONES_JORNADA.values():
                valor = jornada.get(nombre)
                celda = sheet.cell(fila, columnas[nombre])
                celda.value = (
                    datetime.strptime(valor, "%H:%M:%S").time()
                    if isinstance(valor, str) else valor
                ) if valor else None
                celda.number_format = "hh:mm:ss"
            resultado = sheet.cell(fila, columnas["horas_trabajadas"])
            entrada, inicio, fin, salida = (
                f"{get_column_letter(columnas[nombre])}{fila}"
                for nombre in cls.ACCIONES_JORNADA.values()
            )
            resultado.value = (
                f'=IF(OR({entrada}="",{salida}=""),"",'
                f'MOD({salida}-{entrada},1)-IF(OR({inicio}="",{fin}=""),0,'
                f'MOD({fin}-{inicio},1)))'
            )
            resultado.number_format = "[h]:mm"
            jornadas.pop(clave, None)

    # Valida los campos obligatorios y añade el fichaje.
    def guardar_registro(
        self,
        id_empleado,
        nombre_completo,
        departamento,
        accion,
        observaciones="",
        hora_inicio_prevista="",
        hora_fin_prevista="",
        hora_inicio_descanso_prevista="",
        hora_fin_descanso_prevista="",
    ):
        if not isinstance(id_empleado, str) or not id_empleado.strip():
            raise ValueError("Introduce el ID del empleado (por ejemplo, EMP-001).")

        if not isinstance(nombre_completo, str) or not nombre_completo.strip():
            raise ValueError("Introduce el nombre completo del empleado.")

        if not isinstance(departamento, str) or not departamento.strip():
            raise ValueError("Selecciona el departamento del empleado.")

        if not isinstance(accion, str) or not accion.strip():
            raise ValueError("Selecciona el tipo de registro.")

        if not isinstance(observaciones, str):
            raise ValueError("Las observaciones deben ser un texto.")

        if not isinstance(hora_inicio_prevista, str):
            raise ValueError("La hora de inicio prevista debe indicarse como texto en formato HH:MM.")

        if not isinstance(hora_fin_prevista, str):
            raise ValueError("La hora de fin prevista debe indicarse como texto en formato HH:MM.")

        if not isinstance(hora_inicio_descanso_prevista, str):
            raise ValueError("La hora de inicio del descanso debe indicarse como texto en formato HH:MM.")

        if not isinstance(hora_fin_descanso_prevista, str):
            raise ValueError("La hora de fin del descanso debe indicarse como texto en formato HH:MM.")

        # Genera el número de registro y captura la fecha y hora actuales.
        registros = self.consultar_registros()
        ahora = datetime.now()
        validar_fichaje(registros, id_empleado, accion, ahora=ahora)
        ids = [
            r["id_registro"]
            for r in registros
            if isinstance(r.get("id_registro"), int)
            and not isinstance(r.get("id_registro"), bool)
        ]
        nuevo_id = max(ids, default=0) + 1

        registro = {
            "id_registro": nuevo_id,
            "id_empleado": id_empleado.strip(),
            "nombre_completo": nombre_completo.strip(),
            "departamento": departamento.strip(),
            "accion": accion.strip(),
            "fecha": ahora.strftime("%Y-%m-%d"),
            "hora": ahora.strftime("%H:%M:%S"),
            "observaciones": observaciones.strip(),
            "hora_inicio_prevista": hora_inicio_prevista.strip(),
            "hora_fin_prevista": hora_fin_prevista.strip(),
            "hora_inicio_descanso_prevista": hora_inicio_descanso_prevista.strip(),
            "hora_fin_descanso_prevista": hora_fin_descanso_prevista.strip(),
        }

        workbook = load_workbook(self.filename)

        try:
            sheet = workbook["Registros"]
            headers = [
                sheet.cell(row=1, column=col).value
                for col in range(1, sheet.max_column + 1)
            ]
            # Las fórmulas preparadas hasta la fila 2000 no son fichajes.
            columnas_datos = [i + 1 for i, h in enumerate(headers) if h in self.HEADERS[:8]]
            ultima_fila = max(
                (fila for fila in range(2, sheet.max_row + 1)
                 if any(sheet.cell(fila, col).value not in (None, "")
                        and sheet.cell(fila, col).data_type != "f"
                        for col in columnas_datos)),
                default=1,
            )
            nueva_fila = ultima_fila + 1
            for columna, header in enumerate(headers, 1):
                celda = sheet.cell(nueva_fila, columna)
                if header in registro:
                    celda.value = registro[header]
                elif header in ("diferencia_min", "cumplimiento") and celda.value is None and nueva_fila > 2:
                    anterior = sheet.cell(nueva_fila - 1, columna)
                    if anterior.data_type == "f":
                        celda.value = Translator(
                            anterior.value, origin=anterior.coordinate
                        ).translate_formula(celda.coordinate)
            for tabla in sheet.tables.values():
                inicio_col, inicio_fila, fin_col, fin_fila = range_boundaries(tabla.ref)
                if inicio_fila == 1 and inicio_col == 1:
                    tabla.ref = f"A1:{get_column_letter(fin_col)}{max(fin_fila, nueva_fila)}"
                    if tabla.autoFilter is not None:
                        tabla.autoFilter.ref = tabla.ref
            self._actualizar_jornadas(sheet)
            workbook.save(self.filename)
        finally:
            workbook.close()

        return registro

    # Consulta los registros con filtros opcionales de empleado, departamento y acción.
    def consultar_registros(
        self, empleado="", departamento="", accion=""
    ):
        workbook = load_workbook(self.filename, read_only=True)

        try:
            sheet = workbook["Registros"]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if not headers:
                return []

            columnas_datos = [
                indice for indice, nombre in enumerate(headers)
                if nombre in self.HEADERS[:8]
            ]
            registros = [
                dict(zip(headers, row))
                for row in rows
                if any(
                    indice < len(row) and row[indice] not in (None, "")
                    for indice in columnas_datos
                )
            ]

            # openpyxl no calcula fórmulas. Resolver la búsqueda de horarios
            # utilizada por este libro sin alterar las fórmulas del archivo.
            horarios = {}
            if "Departamentos" in workbook.sheetnames:
                for fila in workbook["Departamentos"].iter_rows(
                    min_row=2, max_row=100, max_col=3, values_only=True
                ):
                    if fila[0] is not None:
                        horarios.setdefault(str(fila[0]).casefold(), fila[1:3])

            for registro in registros:
                for indice, campo in enumerate(
                    ("hora_inicio_prevista", "hora_fin_prevista")
                ):
                    valor = registro.get(campo)
                    if isinstance(valor, str) and valor.startswith("="):
                        patron = (
                            r'=IFERROR\(VLOOKUP\(D\d+,Departamentos!'
                            r'\$A\$2:\$D\$100,' + str(indice + 2)
                            + r',FALSE\),""\)'
                        )
                        if re.fullmatch(patron, valor, flags=re.IGNORECASE):
                            horario = horarios.get(
                                str(registro.get("departamento") or "").casefold(),
                                (None, None),
                            )
                            valor = horario[indice]
                        else:
                            # Nunca mostrar código de fórmula como una hora.
                            valor = None
                    if isinstance(valor, (datetime, time)):
                        valor = valor.strftime("%H:%M")
                    elif isinstance(valor, (int, float)) and 0 <= valor < 1:
                        segundos = round(valor * 86400) % 86400
                        valor = f"{segundos // 3600:02d}:{segundos % 3600 // 60:02d}"
                    registro[campo] = valor

            if empleado.strip():
                termino = empleado.strip().lower()
                registros = [
                    r for r in registros
                    if termino in str(r.get("id_empleado") or "").lower()
                    or termino in str(r.get("nombre_completo") or "").lower()
                ]

            if departamento.strip():
                termino = departamento.strip().lower()
                registros = [
                    r for r in registros
                    if termino in str(r.get("departamento") or "").lower()
                ]

            if accion.strip():
                termino = accion.strip().lower()
                registros = [
                    r for r in registros
                    if termino in str(r.get("accion") or "").lower()
                ]

            return registros

        finally:
            workbook.close()

    def consultar_departamentos(self):
        """Devuelve los departamentos de la hoja Departamentos de Excel."""
        workbook = load_workbook(self.filename, read_only=True)

        try:
            if "Departamentos" not in workbook.sheetnames:
                return []

            sheet = workbook["Departamentos"]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if not headers:
                return []

            columna = next(
                (
                    indice
                    for indice, encabezado in enumerate(headers)
                    if str(encabezado or "").strip().lower() == "departamento"
                ),
                None,
            )

            if columna is None:
                return []

            departamentos = []

            for fila in rows:
                if columna < len(fila):
                    valor = fila[columna]

                    if valor is not None and str(valor).strip():
                        departamentos.append(str(valor).strip())

            return list(dict.fromkeys(departamentos))

        finally:
            workbook.close()


    @staticmethod
    def _hora_departamento(valor):
        """Normaliza horas de texto, de Excel y de tipo time a HH:MM."""
        if valor in (None, ""):
            return ""
        if isinstance(valor, (datetime, time)):
            return valor.strftime("%H:%M")
        if isinstance(valor, (int, float)) and not isinstance(valor, bool) and 0 <= valor < 1:
            minutos = round(valor * 1440) % 1440
            return f"{minutos // 60:02d}:{minutos % 60:02d}"
        if isinstance(valor, str):
            for formato in ("%H:%M", "%H:%M:%S"):
                try:
                    return datetime.strptime(valor.strip(), formato).strftime("%H:%M")
                except ValueError:
                    pass
        raise ValueError(f"Hora prevista no válida en la hoja Departamentos: {valor!r}")

    def consultar_horario_departamento(self, departamento):
        """Lee el horario actual del departamento, sin modificar el libro."""
        vacio = {
            "hora_inicio_prevista": "",
            "hora_fin_prevista": "",
            "hora_inicio_descanso_prevista": "",
            "hora_fin_descanso_prevista": "",
        }
        if not departamento.strip():
            return vacio
        workbook = load_workbook(self.filename, read_only=True, data_only=True)
        try:
            if "Departamentos" not in workbook.sheetnames:
                return vacio
            rows = workbook["Departamentos"].iter_rows(values_only=True)
            headers = next(rows, ())
            columnas = {str(h or "").strip().casefold(): i for i, h in enumerate(headers)}
            def columna(*nombres):
                return next((columnas[n] for n in nombres if n in columnas), None)
            dept = columna("departamento", "departamentos")
            inicio = columna("hora entrada", "hora_inicio_prevista", "hora inicio prevista")
            fin = columna("hora salida", "hora_fin_prevista", "hora final prevista", "hora fin prevista")
            inicio_descanso = columna("inicio previsto descanso", "hora_inicio_descanso_prevista")
            fin_descanso = columna("fin previsto descanso", "hora_fin_descanso_prevista")
            if any(i is None for i in (dept, inicio, fin)):
                raise ValueError("La hoja Departamentos debe incluir Departamento, Hora entrada y Hora salida.")
            for row in rows:
                if str(row[dept] or "").strip().casefold() == departamento.strip().casefold():
                    horario = {
                        "hora_inicio_prevista": self._hora_departamento(row[inicio]),
                        "hora_fin_prevista": self._hora_departamento(row[fin]),
                    }
                    for campo, indice in (
                        ("hora_inicio_descanso_prevista", inicio_descanso),
                        ("hora_fin_descanso_prevista", fin_descanso),
                    ):
                        horario[campo] = (
                            self._hora_departamento(row[indice])
                            if indice is not None and row[indice] not in (None, "")
                            else ""
                        )
                    return horario
            return vacio
        finally:
            workbook.close()
