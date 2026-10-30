import openpyxl as pxl
from pathlib import Path
from tkinter import messagebox
from config import  NOMBRE_EXCEL, COLUMNAS_EXCEL

# PATH = Path(__file__).resolve().parent
# RUTA = PATH / NOMBRE_EXCEL

class GestorExcel:
    def __init__(self, ruta, nombre_excel):
        self.ruta = ruta
        self.nombre_excel = nombre_excel

    def inicializar_excel(self):
        print(f"Inicializando Excel en la ruta: {self.ruta}")
        """Inicializa el archivo de Excel y devuelve el libro y la hoja activa."""
        if not self.ruta.exists():
            libro = pxl.Workbook()
            hoja = libro.active
            hoja.title = "Registros"
            hoja.append(COLUMNAS_EXCEL)
            try:
                libro.save(self.rutal)
            except Exception:
                libro.close()
                raise 
            return libro, hoja
        
        libro = pxl.load_workbook(self.ruta)
        hoja = libro.active
        return libro, hoja
        
    def obtener_registros_excel(self):
        """Obtenemos todos los registros del Excel."""
        libro = None
        try:
            libro, hoja = self.inicializar_excel()
            registros = list(hoja.iter_rows(min_row=2, values_only=True))
            return registros
        finally:
            if libro is not None:
                libro.close()
                
    def guardar_excel(self, registro):
        print(f"registro{registro}")
        libro = None
        try:
            libro, hoja = self.inicializar_excel()
            hoja.append(registro)
            libro.save(self.ruta)
        finally:
            if libro is not None:
                libro.close()
                
